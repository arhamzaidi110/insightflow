import csv
import base64
import json
import io
import zipfile
from datetime import date, datetime
from typing import Any

import pandas as pd
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook

from .config import settings
from .analytics import export_csv, run_plan

router = APIRouter()
MAX_UPLOAD_BYTES = settings.max_upload_mb * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 250
MAX_CELLS = 1_500_000
MAX_XLSX_ENTRIES = 5_000
MAX_XLSX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
PREVIEW_ROWS = 10


def _check_extension(filename: str) -> str:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return "csv"
    if lower.endswith(".xlsx"):
        return "xlsx"
    raise HTTPException(status_code=415, detail="Only CSV and XLSX files are supported.")


def _unique_headers(values: list[Any]) -> tuple[list[str], list[str]]:
    headers: list[str] = []
    warnings: list[str] = []
    counts: dict[str, int] = {}
    used: set[str] = set()
    for index, value in enumerate(values, start=1):
        base = "" if value is None else str(value).strip()
        if not base:
            base = f"Column {index}"
            warnings.append(f"Blank header at column {index} was renamed to '{base}'.")
        count = counts.get(base, 0) + 1
        counts[base] = count
        candidate = base if count == 1 else f"{base} ({count})"
        while candidate in used:
            count += 1
            counts[base] = count
            candidate = f"{base} ({count})"
        if count > 1:
            warnings.append(f"Duplicate header '{base}' was renamed to '{candidate}'.")
        used.add(candidate)
        headers.append(candidate)
    if not headers:
        raise HTTPException(status_code=400, detail="The file does not contain a header row.")
    if len(headers) > MAX_COLUMNS:
        raise HTTPException(status_code=413, detail=f"Files may contain at most {MAX_COLUMNS} columns.")
    return headers, warnings


def _infer_type(series: pd.Series) -> str:
    present = series.dropna()
    if present.empty:
        return "unknown"
    if pd.api.types.is_bool_dtype(series.dtype):
        return "boolean"
    if pd.api.types.is_integer_dtype(series.dtype):
        return "integer"
    if pd.api.types.is_float_dtype(series.dtype):
        return "decimal"
    if pd.api.types.is_datetime64_any_dtype(series.dtype):
        return "datetime"
    if pd.api.types.is_object_dtype(series.dtype):
        converted = pd.to_numeric(present, errors="coerce")
        if converted.notna().any() and converted.notna().sum() != len(present):
            return "mixed"
    return "string"


def _json_value(value: Any) -> Any:
    if value is None:
        return None
    try:
        missing = pd.isna(value)
        if isinstance(missing, bool) and missing:
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _profile(frame: pd.DataFrame, warnings: list[str]) -> dict[str, Any]:
    columns = []
    for name in frame.columns:
        series = frame[name]
        blanks = series.map(lambda value: isinstance(value, str) and not value.strip())
        nulls = series.isna()
        present = series[~(blanks | nulls)]
        kind = _infer_type(present)
        numeric = pd.to_numeric(present, errors="coerce")
        date_like = any(token in str(name).lower() for token in ("date", "time", "_at"))
        dates = pd.to_datetime(present, errors="coerce", format="mixed") if date_like else pd.Series(dtype="datetime64[ns]")
        date_count = int(dates.notna().sum()) if date_like else 0
        numeric_count = int(numeric.notna().sum())
        if date_like and date_count == len(present) and len(present):
            kind = "datetime" if dates.astype(str).str.contains(r"\d{2}:\d{2}").any() else "date"
        elif numeric_count == len(present) and len(present) and kind not in ("boolean",):
            kind = "integer" if ((numeric % 1) == 0).all() else "decimal"
        elif numeric_count and numeric_count < len(present):
            kind = "mixed"
        record: dict[str, Any] = {
            "name": str(name), "inferredType": kind,
            "missingCount": int((blanks | nulls).sum()), "missingPercent": round(float((blanks | nulls).mean() * 100), 2),
            "nullCount": int(nulls.sum()), "blankCount": int(blanks.sum()),
            "distinctCount": int(present.nunique()), "suggestions": [], "warnings": [],
        }
        lowered = str(name).lower()
        if any(word in lowered for word in ("id", "key", "code")):
            record["suggestions"].append("Possible identifier column; verify uniqueness before using it as a key.")
        if date_like:
            record["suggestions"].append("Possible date or time column; suggestion based on its name.")
        if any(word in lowered for word in ("price", "cost", "amount", "revenue", "currency", "total")):
            record["suggestions"].append("Possible currency or amount column; verify its units and currency.")
        if kind in ("integer", "decimal") and numeric_count:
            values = numeric.dropna()
            record["numericStats"] = {"min": float(values.min()), "max": float(values.max()), "mean": float(values.mean()), "median": float(values.median()), "sum": float(values.sum()), "standardDeviation": float(values.std()) if len(values) > 1 else None}
        if date_like and date_count:
            valid_dates = dates.dropna()
            record["dateStats"] = {"min": valid_dates.min().isoformat(), "max": valid_dates.max().isoformat(), "coverageCount": date_count, "coveragePercent": round(date_count / max(1, len(present)) * 100, 2)}
        if kind in ("string", "boolean", "mixed"):
            record["topValues"] = [{"value": str(value), "count": int(count)} for value, count in present.astype(str).value_counts().head(10).items()]
        if kind == "mixed": record["warnings"].append("Values do not share one consistent inferred type.")
        if record["missingPercent"] > 50: record["severity"] = "high"; record["explanation"] = "More than half of this column is blank or null."
        elif kind == "mixed": record["severity"] = "medium"; record["explanation"] = "Some values do not match the dominant inferred type."
        elif record["missingCount"]: record["severity"] = "low"; record["explanation"] = "This column contains blank or null values."
        else: record["severity"] = "info"; record["explanation"] = "No missing values or mixed types were detected."
        columns.append(record)
    return {"rowCount": int(len(frame)), "columnCount": int(len(frame.columns)), "duplicateRowCount": int(frame.duplicated().sum()), "columns": columns, "warnings": warnings}


def _result(frame: pd.DataFrame, warnings: list[str], sheets: list[str], selected_sheet: str | None, include_frame: bool = False) -> dict[str, Any]:
    if frame.empty:
        raise HTTPException(status_code=400, detail="The file has headers but no data rows.")
    if len(frame) > MAX_ROWS:
        raise HTTPException(status_code=413, detail=f"Files may contain at most {MAX_ROWS} data rows.")
    if len(frame.columns) * len(frame) > MAX_CELLS:
        raise HTTPException(status_code=413, detail=f"Files may contain at most {MAX_CELLS} data cells.")

    preview_frame = frame.head(PREVIEW_ROWS).astype(object)
    preview = [
        {str(column): _json_value(value) for column, value in row.items()}
        for row in preview_frame.where(pd.notna(preview_frame), None).to_dict(orient="records")
    ]
    result = {
        "rowCount": int(len(frame)),
        "columnCount": int(len(frame.columns)),
        "columns": [{"name": str(column), "inferredType": _infer_type(frame[column])} for column in frame.columns],
        "preview": preview,
        "warnings": warnings,
        "sheets": sheets,
        "selectedSheet": selected_sheet,
        "profile": _profile(frame, warnings),
    }
    if include_frame:
        result["_frame"] = frame
    return result


def _parse_csv(data: bytes, include_frame: bool = False) -> dict[str, Any]:
    warnings: list[str] = []
    try:
        text = data.decode("utf-8-sig")
        encoding = "utf-8-sig"
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
        encoding = "cp1252"
        warnings.append("The file was decoded with Windows-1252 because it is not valid UTF-8.")
    if "\x00" in text:
        raise HTTPException(status_code=400, detail="The CSV contains binary or null-byte content.")
    try:
        sample = text[:65_536]
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
        warnings.append("The delimiter could not be inferred; comma was used.")
    try:
        header_row = next(csv.reader(io.StringIO(text), delimiter=delimiter))
    except (StopIteration, csv.Error):
        raise HTTPException(status_code=400, detail="The CSV is empty or malformed.") from None
    headers, header_warnings = _unique_headers(header_row)
    warnings.extend(header_warnings)
    max_data_rows = min(MAX_ROWS, MAX_CELLS // len(headers))
    try:
        frame = pd.read_csv(
            io.BytesIO(data),
            sep=delimiter,
            header=0,
            names=headers,
            encoding=encoding,
            engine="python",
            on_bad_lines="error",
            nrows=max_data_rows + 1,
            skip_blank_lines=False,
            keep_default_na=False,
            na_values=["NA", "N/A", "NULL", "null"],
        )
    except (pd.errors.EmptyDataError, pd.errors.ParserError, UnicodeError, ValueError) as error:
        raise HTTPException(status_code=400, detail=f"The CSV could not be parsed: {type(error).__name__}.") from None
    if len(frame) > max_data_rows:
        reason = f"{MAX_ROWS} data rows" if MAX_ROWS <= MAX_CELLS // len(headers) else f"{MAX_CELLS} data cells"
        raise HTTPException(status_code=413, detail=f"Files may contain at most {reason}.")
    blank_mask = frame.apply(lambda column: column.map(lambda value: pd.isna(value) or (isinstance(value, str) and not value.strip()))).all(axis=1)
    blank_rows = int(blank_mask.sum())
    if blank_rows:
        frame = frame.loc[~blank_mask].reset_index(drop=True)
        warnings.append(f"{blank_rows} completely blank CSV row(s) were ignored.")
    return _result(frame, warnings, [], None, include_frame)


def _check_xlsx_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise HTTPException(status_code=413, detail="The workbook contains too many internal files.")
            expanded_size = sum(entry.file_size for entry in entries)
            if expanded_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise HTTPException(status_code=413, detail="The expanded workbook is larger than the allowed limit.")
            if any(entry.compress_size == 0 and entry.file_size > 0 for entry in entries):
                raise HTTPException(status_code=400, detail="The workbook archive is malformed.")
            if any(
                entry.compress_size > 0 and entry.file_size / entry.compress_size > MAX_COMPRESSION_RATIO
                for entry in entries
            ):
                raise HTTPException(status_code=413, detail="The workbook compression ratio exceeds the allowed limit.")
            if not any(entry.filename == "[Content_Types].xml" for entry in entries):
                raise HTTPException(status_code=400, detail="The file is not a valid XLSX workbook.")
    except zipfile.BadZipFile:
        raise HTTPException(status_code=400, detail="The XLSX workbook is malformed.") from None


def _parse_xlsx(data: bytes, selected_sheet: str | None, include_frame: bool = False) -> dict[str, Any]:
    _check_xlsx_archive(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except Exception as error:
        raise HTTPException(status_code=400, detail=f"The XLSX workbook could not be read: {type(error).__name__}.") from None
    try:
        sheets = list(workbook.sheetnames)
        if not sheets:
            raise HTTPException(status_code=400, detail="The workbook does not contain any worksheets.")
        sheet_name = selected_sheet or sheets[0]
        if sheet_name not in sheets:
            raise HTTPException(status_code=400, detail="The selected worksheet does not exist in this workbook.")
        worksheet = workbook[sheet_name]
        if worksheet.max_column and worksheet.max_column > MAX_COLUMNS:
            raise HTTPException(status_code=413, detail=f"Worksheets may contain at most {MAX_COLUMNS} columns.")
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_row = list(next(rows))
        except StopIteration:
            raise HTTPException(status_code=400, detail="The selected worksheet is empty.") from None
        while header_row and header_row[-1] is None:
            header_row.pop()
        headers, warnings = _unique_headers(header_row)
        values: list[list[Any]] = []
        ignored_blank_rows = 0
        for row in rows:
            if any(value is not None for value in row[len(headers) :]):
                raise HTTPException(status_code=400, detail="The worksheet contains values beyond its header columns.")
            row_values = list(row[: len(headers)])
            if not any(value is not None for value in row_values):
                ignored_blank_rows += 1
                continue
            values.append(row_values)
            if len(values) > MAX_ROWS:
                raise HTTPException(status_code=413, detail=f"Worksheets may contain at most {MAX_ROWS} data rows.")
            if len(values) * len(headers) > MAX_CELLS:
                raise HTTPException(status_code=413, detail=f"Worksheets may contain at most {MAX_CELLS} data cells.")
        if ignored_blank_rows:
            warnings.append(f"{ignored_blank_rows} completely blank worksheet row(s) were ignored.")
        frame = pd.DataFrame(values, columns=headers)
        return _result(frame, warnings, sheets, sheet_name, include_frame)
    finally:
        workbook.close()


@router.post("/preview")
async def preview_upload(
    file: UploadFile = File(...),
    sheet_name: str | None = Form(default=None),
) -> dict[str, Any]:
    file_kind = _check_extension(file.filename or "")
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    await file.close()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"Files may not exceed {settings.max_upload_mb} MB.")
    if not data:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if file_kind == "xlsx":
        return _parse_xlsx(data, sheet_name)
    return _parse_csv(data)


def _clean_frame(frame: pd.DataFrame, operation: dict[str, Any], apply: bool = False) -> tuple[pd.DataFrame, list[str]]:
    action = operation.get("action")
    warnings: list[str] = []
    if action == "trim_whitespace":
        for column in frame.select_dtypes(include="object").columns:
            frame[column] = frame[column].map(lambda value: value.strip() if isinstance(value, str) else value)
    elif action == "standardize_blanks":
        frame = frame.replace(r"^\s*$", pd.NA, regex=True)
    elif action == "remove_duplicates":
        frame = frame.drop_duplicates().reset_index(drop=True)
    elif action == "drop_columns":
        columns = operation.get("columns")
        if not isinstance(columns, list) or not columns or any(column not in frame.columns for column in columns):
            raise HTTPException(status_code=422, detail="Select one or more existing columns to remove.")
        if len(columns) >= len(frame.columns):
            raise HTTPException(status_code=422, detail="At least one column must remain.")
        frame = frame.drop(columns=columns)
    elif action == "rename_column":
        old, new = operation.get("column"), operation.get("newName")
        if old not in frame.columns or not isinstance(new, str) or not new.strip() or len(new.strip()) > 120:
            raise HTTPException(status_code=422, detail="Choose a column and provide a valid name (1–120 characters).")
        new = new.strip()
        if new != old and new in frame.columns:
            raise HTTPException(status_code=422, detail="A column with that name already exists.")
        frame = frame.rename(columns={old: new})
    elif action in ("convert_numeric", "convert_date"):
        column = operation.get("column")
        if column not in frame.columns:
            raise HTTPException(status_code=422, detail="Choose an existing column.")
        source = frame[column].replace(r"^\s*$", pd.NA, regex=True)
        converted = pd.to_numeric(source, errors="coerce") if action == "convert_numeric" else pd.to_datetime(source, errors="coerce", format="mixed")
        failures = source.notna() & converted.isna()
        failure_count = int(failures.sum())
        if failure_count and not operation.get("coerceFailures", False):
            sample_failures = source.loc[failures].astype(str).head(5).tolist()
            if apply: raise HTTPException(status_code=422, detail="Some values fail conversion. Preview them, then explicitly choose to replace failures with missing values.")
            warnings.append(f"{int(source.notna().sum())-failure_count} value(s) will convert; {failure_count} will fail. Review failures {sample_failures}, then choose whether to replace them with missing values.")
        else:
            frame[column] = converted
    elif action == "fill_missing":
        column, method = operation.get("column"), operation.get("method")
        if column not in frame.columns:
            raise HTTPException(status_code=422, detail="Choose an existing column.")
        series = frame[column]
        if method in ("mean", "median"):
            numeric = pd.to_numeric(series, errors="coerce")
            if numeric.notna().sum() == 0:
                raise HTTPException(status_code=422, detail="Mean or median requires numeric values in the selected column.")
            fill_value = numeric.mean() if method == "mean" else numeric.median()
        elif method == "mode":
            modes = series.dropna().mode()
            if modes.empty:
                raise HTTPException(status_code=422, detail="Mode requires at least one non-missing value.")
            fill_value = modes.iloc[0]
        elif method == "constant":
            fill_value = operation.get("value")
            if fill_value is None or isinstance(fill_value, (dict, list)):
                raise HTTPException(status_code=422, detail="Enter a text or numeric constant.")
            if pd.api.types.is_numeric_dtype(series.dtype):
                try: fill_value = float(fill_value)
                except (ValueError, TypeError): raise HTTPException(status_code=422, detail="The selected column requires a numeric constant.") from None
        else:
            raise HTTPException(status_code=422, detail="Choose mean, median, mode, or constant.")
        frame[column] = series.fillna(fill_value)
    elif action == "filter_rows":
        column, operator, value = operation.get("column"), operation.get("operator"), operation.get("value")
        if column not in frame.columns:
            raise HTTPException(status_code=422, detail="Choose an existing column.")
        series = frame[column]
        try:
            if operator == "eq": mask = series == value
            elif operator == "ne": mask = series != value
            elif operator == "contains": mask = series.astype(str).str.contains(str(value), regex=False, na=False)
            elif operator in ("gt", "gte", "lt", "lte"):
                numeric = pd.to_numeric(series, errors="coerce")
                non_numeric = int((series.notna() & numeric.isna()).sum())
                if numeric.notna().sum() == 0: raise ValueError("not numeric")
                if non_numeric: warnings.append(f"{non_numeric} non-numeric value(s) did not match this numeric filter.")
                number = float(value)
                if operator == "gt": mask = numeric > number
                elif operator == "gte": mask = numeric >= number
                elif operator == "lt": mask = numeric < number
                else: mask = numeric <= number
            else: raise ValueError("unsupported operator")
        except (TypeError, ValueError):
            raise HTTPException(status_code=422, detail="Choose a supported comparison and a compatible value.") from None
        frame = frame.loc[mask].reset_index(drop=True)
        if frame.empty:
            raise HTTPException(status_code=422, detail="This filter would remove every row.")
    else:
        raise HTTPException(status_code=422, detail="Choose a supported cleaning action.")
    return frame, warnings


@router.post("/clean")
async def clean_upload(file: UploadFile = File(...), operation: str = Form(...), sheet_name: str | None = Form(default=None), apply: bool = Form(default=False)) -> dict[str, Any]:
    kind = _check_extension(file.filename or "")
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    await file.close()
    if len(data) > MAX_UPLOAD_BYTES: raise HTTPException(status_code=413, detail=f"Files may not exceed {settings.max_upload_mb} MB.")
    try: request = json.loads(operation)
    except ValueError: raise HTTPException(status_code=422, detail="Invalid cleaning operation.") from None
    if not isinstance(request, dict): raise HTTPException(status_code=422, detail="Invalid cleaning operation.")
    parsed = _parse_xlsx(data, sheet_name, True) if kind == "xlsx" else _parse_csv(data, True)
    original = parsed.pop("_frame")
    before_rows, before_columns = len(original), len(original.columns)
    cleaned, warnings = _clean_frame(original.copy(), request, apply)
    if len(cleaned) * len(cleaned.columns) > MAX_CELLS: raise HTTPException(status_code=413, detail="The cleaned data exceeds the supported cell limit.")
    report = _result(cleaned, [], parsed["sheets"], parsed["selectedSheet"])
    csv_text = cleaned.to_csv(index=False, na_rep="")
    return {"before": {"rowCount": before_rows, "columnCount": before_columns}, "after": {"rowCount": len(cleaned), "columnCount": len(cleaned.columns)},
            "affectedRows": abs(before_rows-len(cleaned)), "affectedColumns": abs(before_columns-len(cleaned.columns)), "warnings": warnings,
            "preview": report["preview"], "columns": report["columns"], "profile": report["profile"], "contentBase64": base64.b64encode(csv_text.encode("utf-8-sig")).decode("ascii")}


@router.post("/analyze")
async def analyze_upload(file: UploadFile = File(...), plan: str = Form(...), sheet_name: str | None = Form(default=None)) -> dict[str, Any]:
    kind = _check_extension(file.filename or "")
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    await file.close()
    if len(data) > MAX_UPLOAD_BYTES: raise HTTPException(status_code=413, detail=f"Files may not exceed {settings.max_upload_mb} MB.")
    try: request = json.loads(plan)
    except ValueError: raise HTTPException(status_code=422, detail="The analysis plan must be valid JSON.") from None
    parsed = _parse_xlsx(data, sheet_name, True) if kind == "xlsx" else _parse_csv(data, True)
    frame = parsed.pop("_frame")
    return run_plan(frame, request)


@router.post("/export")
async def export_upload(file: UploadFile = File(...), plan: str = Form(...), sheet_name: str | None = Form(default=None)) -> dict[str, str]:
    kind = _check_extension(file.filename or "")
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    await file.close()
    if len(data) > MAX_UPLOAD_BYTES: raise HTTPException(status_code=413, detail=f"Files may not exceed {settings.max_upload_mb} MB.")
    try: request = json.loads(plan)
    except ValueError: raise HTTPException(status_code=422, detail="The export plan must be valid JSON.") from None
    if not isinstance(request, dict): raise HTTPException(status_code=422, detail="The export plan must be a JSON object.")
    parsed = _parse_xlsx(data, sheet_name, True) if kind == "xlsx" else _parse_csv(data, True)
    frame = parsed.pop("_frame")
    return {"contentBase64": base64.b64encode(export_csv(frame, request)).decode("ascii")}
